"""
Stream-lite — ANOVA Builder
===========================
Upload an Excel/CSV master chart, pick a sheet, fix variable types, choose one or
more outcome (numerical) variables and one or more grouping (categorical)
variables, and get:

  1. Assumption checks   - Shapiro-Wilk normality per group + Levene (Brown-Forsythe)
  2. Summary table       - mean ± SD, median (IQR) or both (or chosen from normality),
                           with test name, test statistic, p-value and effect size
  3. Post-hoc pairwise   - difference, CI, test statistic and p-value for every pair

Run with:
    pip install streamlit pandas numpy scipy openpyxl python-docx
    streamlit run anova_app.py

How test selection works
------------------------
Parametric      -> One-way ANOVA (equal variances) or Welch's ANOVA (unequal variances)
                   Post hoc: Tukey HSD, Games-Howell or Bonferroni t-tests
Non-parametric  -> Kruskal-Wallis H
                   Post hoc: pairwise Mann-Whitney U (Holm / Bonferroni / BH / none)
                   with Hodges-Lehmann median difference and CI
Auto            -> parametric if every group passes Shapiro-Wilk (p > alpha),
                   otherwise non-parametric. You can force either one.
"""

import html as _html
import io
import itertools
import json
import re

import numpy as np
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from scipy import stats
from scipy.stats import studentized_range
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

st.set_page_config(page_title="Stream-lite · ANOVA Builder", layout="wide")

P_COLS = {"p-value", "p (raw)", "p (adj.)"}

# --------------------------------------------------------------------------
# Generic helpers
# --------------------------------------------------------------------------

def detect_type(series: pd.Series):
    """Guess whether a column is numerical or categorical."""
    non_missing = series.dropna()
    non_missing = non_missing[non_missing.astype(str).str.strip() != ""]
    n = len(non_missing)
    if n == 0:
        return "categorical", 0, 0
    numeric_ratio = pd.to_numeric(non_missing, errors="coerce").notna().mean()
    unique_n = non_missing.astype(str).str.strip().nunique()
    if numeric_ratio >= 0.9 and unique_n > 10:
        return "numerical", n, unique_n
    return "categorical", n, unique_n


def effective_type(meta):
    return meta["detected"] if meta["type"] == "auto" else meta["type"]


def fmt_p(p):
    if p is None or (isinstance(p, float) and np.isnan(p)):
        return "—"
    return "<0.001" if p < 0.001 else f"{p:.3f}"


def fmt_num(x, d=2):
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return "—"
    return f"{x:.{d}f}"


def is_sig(s, alpha):
    s = str(s)
    if s.startswith("<"):
        return True
    try:
        return float(s) < alpha
    except ValueError:
        return False


def padjust(p, method):
    p = np.asarray(p, float)
    m = len(p)
    if method == "none" or m <= 1:
        return p.copy()
    if method == "bonferroni":
        return np.minimum(1, p * m)
    order = np.argsort(p)
    adj = np.empty(m)
    if method == "holm":
        running = 0.0
        for rank, i in enumerate(order):
            running = max(running, (m - rank) * p[i])
            adj[i] = min(1.0, running)
        return adj
    if method == "bh":
        running = 1.0
        for rank in range(m - 1, -1, -1):
            i = order[rank]
            running = min(running, p[i] * m / (rank + 1))
            adj[i] = running
        return adj
    raise ValueError(method)


# --------------------------------------------------------------------------
# Assumption checks
# --------------------------------------------------------------------------

def normality_test(arr):
    """Shapiro-Wilk (n<=5000) or D'Agostino-Pearson (n>5000). Returns (stat, p)
    or (None, None) when it cannot be computed (n<3 or constant data)."""
    arr = np.asarray(arr, float)
    if len(arr) < 3 or np.ptp(arr) == 0:
        return None, None
    try:
        if len(arr) > 5000:
            s, p = stats.normaltest(arr)
        else:
            s, p = stats.shapiro(arr)
        return float(s), float(p)
    except Exception:
        return None, None


# --------------------------------------------------------------------------
# Omnibus tests
# --------------------------------------------------------------------------

def eta_squared(groups):
    allv = np.concatenate(groups)
    gm = allv.mean()
    ssb = sum(len(g) * (g.mean() - gm) ** 2 for g in groups)
    sst = ((allv - gm) ** 2).sum()
    return float(ssb / sst) if sst > 0 else np.nan


def classic_anova(groups):
    res = stats.f_oneway(*groups)
    k, N = len(groups), sum(len(g) for g in groups)
    return {"name": "One-way ANOVA", "stat": f"F={res.statistic:.2f}, df={k-1},{N-k}",
            "p": float(res.pvalue), "effect": f"\u03B7\u00B2={eta_squared(groups):.3f}"}


def welch_anova(groups):
    k = len(groups)
    n = np.array([len(g) for g in groups], float)
    m = np.array([g.mean() for g in groups])
    v = np.array([g.var(ddof=1) for g in groups])
    w = n / v
    W = w.sum()
    mw = (w * m).sum() / W
    a = (w * (m - mw) ** 2).sum() / (k - 1)
    t = (((1 - w / W) ** 2) / (n - 1)).sum()
    b = 1 + 2 * (k - 2) / (k ** 2 - 1) * t
    F = a / b
    df2 = (k ** 2 - 1) / (3 * t)
    p = float(stats.f.sf(F, k - 1, df2))
    return {"name": "Welch's ANOVA", "stat": f"F={F:.2f}, df={k-1},{df2:.1f}",
            "p": p, "effect": f"\u03B7\u00B2={eta_squared(groups):.3f}"}


def kruskal(groups):
    res = stats.kruskal(*groups)
    N = sum(len(g) for g in groups)
    eps = res.statistic / (N - 1) if N > 1 else np.nan
    return {"name": "Kruskal\u2013Wallis H", "stat": f"H={res.statistic:.2f}, df={len(groups)-1}",
            "p": float(res.pvalue), "effect": f"\u03B5\u00B2={eps:.3f}"}


# --------------------------------------------------------------------------
# Post-hoc tests
# --------------------------------------------------------------------------

def posthoc_parametric(names, groups, method, alpha, d):
    """method: 'tukey' | 'gameshowell' | 'bonferroni'. Returns list of row dicts."""
    k = len(groups)
    n = np.array([len(g) for g in groups], float)
    m = np.array([g.mean() for g in groups])
    v = np.array([g.var(ddof=1) for g in groups])
    N = n.sum()
    dfw = N - k
    mse = ((n - 1) * v).sum() / dfw
    pairs = list(itertools.combinations(range(k), 2))
    npairs = len(pairs)
    rows = []
    for i, j in pairs:
        diff = m[i] - m[j]
        if method == "tukey":
            se = np.sqrt(mse / 2 * (1 / n[i] + 1 / n[j]))
            q = abs(diff) / se if se > 0 else np.nan
            p = float(studentized_range.sf(q, k, dfw))
            crit = studentized_range.ppf(1 - alpha, k, dfw)
            lo, hi = diff - crit * se, diff + crit * se
            stat = f"q={q:.2f}, df={dfw:.0f}"
            label = "Tukey HSD"
        elif method == "gameshowell":
            s2 = v[i] / n[i] + v[j] / n[j]
            se = np.sqrt(s2 / 2)
            df = s2 ** 2 / ((v[i] / n[i]) ** 2 / (n[i] - 1) + (v[j] / n[j]) ** 2 / (n[j] - 1))
            q = abs(diff) / se if se > 0 else np.nan
            p = float(studentized_range.sf(q, k, df))
            crit = studentized_range.ppf(1 - alpha, k, df)
            lo, hi = diff - crit * se, diff + crit * se
            stat = f"q={q:.2f}, df={df:.1f}"
            label = "Games-Howell"
        else:  # bonferroni pooled-variance t-tests
            se = np.sqrt(mse * (1 / n[i] + 1 / n[j]))
            t = diff / se if se > 0 else np.nan
            p_raw = 2 * stats.t.sf(abs(t), dfw)
            p = float(min(1.0, p_raw * npairs))
            crit = stats.t.ppf(1 - alpha / (2 * npairs), dfw)
            lo, hi = diff - crit * se, diff + crit * se
            stat = f"t={t:.2f}, df={dfw:.0f}"
            label = "Bonferroni t-test"
        rows.append({"Comparison": f"{names[i]} \u2212 {names[j]}", "Method": label,
                     "Difference": fmt_num(diff, d), "CI": f"{fmt_num(lo, d)} to {fmt_num(hi, d)}",
                     "Statistic": stat, "p (raw)": "—", "p (adj.)": fmt_p(p), "_p": p})
    return rows


def hodges_lehmann(a, b, alpha):
    """Median of all pairwise differences a-b with a distribution-free CI
    (normal approximation to the Mann-Whitney U distribution)."""
    diffs = np.sort((a[:, None] - b[None, :]).ravel())
    n1n2 = len(diffs)
    est = float(np.median(diffs))
    z = stats.norm.ppf(1 - alpha / 2)
    c = n1n2 / 2 - z * np.sqrt(n1n2 * (len(a) + len(b) + 1) / 12)
    c = int(np.floor(c))
    lo = diffs[max(c, 0)]
    hi = diffs[min(n1n2 - 1 - c, n1n2 - 1)] if c >= 0 else diffs[-1]
    return est, float(lo), float(hi)


def posthoc_nonparametric(names, groups, padj, alpha, d):
    k = len(groups)
    pairs = list(itertools.combinations(range(k), 2))
    npairs = len(pairs)
    # CI is Bonferroni-widened whenever a familywise correction (Holm/Bonferroni) is chosen
    ci_alpha = alpha / npairs if padj in ("bonferroni", "holm") else alpha
    raw, parts = [], []
    for i, j in pairs:
        a, b = groups[i], groups[j]
        res = stats.mannwhitneyu(a, b, alternative="two-sided")
        est, lo, hi = hodges_lehmann(a, b, ci_alpha)
        raw.append(float(res.pvalue))
        parts.append((i, j, res.statistic, est, lo, hi))
    adj = padjust(raw, padj)
    label = {"holm": "Mann-Whitney (Holm)", "bonferroni": "Mann-Whitney (Bonferroni)",
             "bh": "Mann-Whitney (BH)", "none": "Mann-Whitney (unadjusted)"}[padj]
    rows = []
    for (i, j, U, est, lo, hi), pr, pa in zip(parts, raw, adj):
        rows.append({"Comparison": f"{names[i]} \u2212 {names[j]}", "Method": label,
                     "Difference": fmt_num(est, d), "CI": f"{fmt_num(lo, d)} to {fmt_num(hi, d)}",
                     "Statistic": f"U={U:.1f}", "p (raw)": fmt_p(pr),
                     "p (adj.)": fmt_p(float(pa)), "_p": float(pa)})
    return rows


# --------------------------------------------------------------------------
# Per outcome x grouping analysis
# --------------------------------------------------------------------------

def get_groups(df, outcome, gcol):
    sub = pd.DataFrame({
        "y": pd.to_numeric(df[outcome], errors="coerce"),
        "g": df[gcol].astype(str).str.strip(),
    })
    sub = sub.dropna()
    sub = sub[~sub["g"].isin(["", "nan", "NaN", "None"])]
    out = {}
    dropped = []
    for lvl, grp in sub.groupby("g"):
        if len(grp) >= 2:
            out[lvl] = grp["y"].to_numpy(float)
        else:
            dropped.append(lvl)
    return out, dropped


def summary_cell(values, mode, use_param, d):
    parts = []
    if mode in ("mean_sd", "both") or (mode == "auto" and use_param):
        parts.append(f"{fmt_num(np.mean(values), d)} \u00B1 {fmt_num(np.std(values, ddof=1), d)}")
    if mode in ("median_iqr", "both") or (mode == "auto" and not use_param):
        q1, med, q3 = np.percentile(values, [25, 50, 75])
        parts.append(f"{fmt_num(med, d)} ({fmt_num(q1, d)}\u2013{fmt_num(q3, d)})")
    return "; ".join(parts)


def outcome_label(col, mode, use_param):
    if mode == "mean_sd" or (mode == "auto" and use_param):
        return f"{col}, mean \u00B1 SD"
    if mode == "median_iqr" or (mode == "auto" and not use_param):
        return f"{col}, median (IQR)"
    return f"{col}, mean \u00B1 SD; median (IQR)"


def analyze(df, outcome, gcol, cfg):
    """Returns dict(summary_cells, test, assumptions, posthoc, use_param, msg) or None + message."""
    alpha, d = cfg["alpha"], cfg["decimals"]
    groups, dropped = get_groups(df, outcome, gcol)
    if len(groups) < 2:
        return None, f"{outcome} by {gcol}: fewer than 2 groups with n\u22652 \u2014 skipped."
    names = sorted(groups)
    arrs = [groups[nm] for nm in names]
    k = len(arrs)

    assumptions, normal_flags = [], []
    for nm, a in zip(names, arrs):
        W, p = normality_test(a)
        ok = (p is not None) and p > alpha
        normal_flags.append(ok)
        assumptions.append({
            "Outcome": outcome, "Check": "Normality (Shapiro-Wilk)" if len(a) <= 5000 else "Normality (D'Agostino-Pearson)",
            "Group": nm, "n": len(a), "Statistic": "—" if W is None else f"{W:.3f}",
            "p-value": fmt_p(p), "Result": "Normal" if ok else ("Not normal" if p is not None else "n/a (n<3 or constant)"),
        })
    try:
        lev_stat, lev_p = stats.levene(*arrs, center="median")
        lev_p = float(lev_p)
        eq_var = lev_p > alpha
        assumptions.append({
            "Outcome": outcome, "Check": "Equal variances (Levene, median-centred)", "Group": "All groups",
            "n": sum(len(a) for a in arrs), "Statistic": f"W={lev_stat:.3f}", "p-value": fmt_p(lev_p),
            "Result": "Equal" if eq_var else "Unequal",
        })
    except Exception:
        lev_p, eq_var = np.nan, True

    if cfg["test_mode"] == "parametric":
        use_param = True
    elif cfg["test_mode"] == "nonparametric":
        use_param = False
    else:
        use_param = all(normal_flags)

    try:
        if use_param:
            welch = cfg["variance"] == "welch" or (cfg["variance"] == "auto" and not eq_var)
            if welch and min(a.var(ddof=1) for a in arrs) == 0:
                welch = False  # Welch needs non-zero variance in every group
            test = welch_anova(arrs) if welch else classic_anova(arrs)
        else:
            welch = False
            test = kruskal(arrs)
    except Exception as e:
        return None, f"{outcome} by {gcol}: test could not be computed ({e})."

    posthoc = []
    if k >= 3 and (not cfg["posthoc_only_sig"] or test["p"] < alpha):
        try:
            if use_param:
                m = cfg["posthoc_param"]
                if m == "auto":
                    m = "gameshowell" if welch else "tukey"
                posthoc = posthoc_parametric(names, arrs, m, alpha, d)
            else:
                posthoc = posthoc_nonparametric(names, arrs, cfg["posthoc_np"], alpha, d)
        except Exception as e:
            return None, f"{outcome} by {gcol}: post-hoc failed ({e})."
        for r in posthoc:
            r["Outcome"] = outcome

    cells = {nm: summary_cell(a, cfg["display"], use_param, d) for nm, a in zip(names, arrs)}
    notes = []
    if dropped:
        notes.append(f"{outcome} by {gcol}: groups with n<2 excluded ({', '.join(dropped)}).")
    return {"cells": cells, "test": test, "assumptions": assumptions, "posthoc": posthoc,
            "use_param": use_param, "k": k, "label": outcome_label(outcome, cfg["display"], use_param)}, notes


# --------------------------------------------------------------------------
# Rendering / export
# --------------------------------------------------------------------------

def render_html(df, alpha):
    css = ("<style>.sl-t{border-collapse:collapse;width:100%;font-size:0.9rem;"
           "border-top:2px solid currentColor;border-bottom:2px solid currentColor}"
           ".sl-t th{border-bottom:1px solid currentColor;text-align:left;padding:4px 8px}"
           ".sl-t td{padding:4px 8px}</style>")
    head = "".join(f"<th>{_html.escape(str(c))}</th>" for c in df.columns)
    body = ""
    for _, row in df.iterrows():
        tds = ""
        for c in df.columns:
            v = _html.escape(str(row[c]))
            if c in P_COLS and is_sig(row[c], alpha):
                v = f"<b>{v}</b>"
            tds += f"<td>{v}</td>"
        body += f"<tr>{tds}</tr>"
    return f"{css}<table class='sl-t'><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def copy_button(df, label="Copy table"):
    """Button that copies df to the clipboard as tab-separated text (with header),
    which pastes straight into Excel cells."""
    tsv = df.astype(str).to_csv(sep="\t", index=False)
    payload = json.dumps(tsv)
    components.html(f"""
    <button id="b" style="padding:5px 12px;border:1px solid #888;border-radius:6px;
      background:transparent;color:inherit;cursor:pointer;font-size:13px;font-family:sans-serif">
      \U0001F4CB {label}</button>
    <script>
    const text = {payload};
    const btn = document.getElementById('b');
    const done = ok => {{ const o = btn.innerHTML; btn.innerHTML = ok ? '\u2705 Copied!' : '\u274C Copy failed';
                         setTimeout(() => btn.innerHTML = o, 1500); }};
    function fallback() {{
      const t = document.createElement('textarea'); t.value = text;
      document.body.appendChild(t); t.select();
      let ok = false; try {{ ok = document.execCommand('copy'); }} catch (e) {{}}
      document.body.removeChild(t); done(ok);
    }}
    btn.onclick = () => {{
      if (navigator.clipboard && window.isSecureContext) {{
        navigator.clipboard.writeText(text).then(() => done(true), fallback);
      }} else fallback();
    }};
    </script>""", height=42)


def build_excel(results):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for i, (g, r) in enumerate(results.items(), 1):
            tag = re.sub(r"[\[\]:*?/\\]", "_", g)[:22]
            r["summary"].to_excel(xw, sheet_name=f"S{i}_{tag}"[:31], index=False)
            r["assumptions"].to_excel(xw, sheet_name=f"A{i}_{tag}"[:31], index=False)
            if not r["posthoc"].empty:
                r["posthoc"].to_excel(xw, sheet_name=f"P{i}_{tag}"[:31], index=False)
    buf.seek(0)
    return buf.getvalue()


def _docx_table(doc, df, alpha):
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Table Grid"
    for i, c in enumerate(df.columns):
        t.rows[0].cells[i].text = str(c)
        for run in t.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for i, c in enumerate(df.columns):
            cells[i].text = str(row[c])
            for p in cells[i].paragraphs:
                if i > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(8)
                    if c in P_COLS and is_sig(row[c], alpha):
                        run.bold = True


def build_docx(results, alpha, footnotes):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for g, r in results.items():
        doc.add_heading(f"Comparison by {g}", level=2)
        doc.add_paragraph("Summary").runs[0].bold = True
        _docx_table(doc, r["summary"], alpha)
        if not r["posthoc"].empty:
            doc.add_paragraph()
            doc.add_paragraph("Post-hoc pairwise comparisons").runs[0].bold = True
            _docx_table(doc, r["posthoc"], alpha)
        doc.add_paragraph()
        doc.add_paragraph("Assumption checks").runs[0].bold = True
        _docx_table(doc, r["assumptions"], alpha)
    doc.add_paragraph()
    for f in footnotes:
        p = doc.add_paragraph(f)
        for run in p.runs:
            run.italic = True
            run.font.size = Pt(8)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


# --------------------------------------------------------------------------
# Streamlit UI
# --------------------------------------------------------------------------

st.title("Stream-lite · ANOVA Builder")
st.caption(
    "Upload a master chart, pick the outcome(s) and grouping variable(s), and get descriptive "
    "statistics, parametric or non-parametric one-way ANOVA, and post-hoc pairwise comparisons "
    "with confidence intervals, test statistics and p-values."
)

st.markdown("### 1. Upload master chart")
uploaded = st.file_uploader("Excel (.xlsx/.xls) or CSV. First row must be column headers.",
                            type=["xlsx", "xls", "csv"])

if uploaded is None:
    st.info("Upload a file to get started. Nothing leaves your machine \u2014 the app runs locally.")
    st.stop()

sheet_name = None
try:
    if not uploaded.name.lower().endswith(".csv"):
        xls = pd.ExcelFile(uploaded)
        sheet_names = xls.sheet_names
        sheet_name = st.selectbox("Select sheet", options=sheet_names) if len(sheet_names) > 1 else sheet_names[0]
        df = pd.read_excel(xls, sheet_name=sheet_name)
    else:
        df = pd.read_csv(uploaded)
except Exception as e:
    st.error(f"Could not read that file: {e}")
    st.stop()

df.columns = [str(c) for c in df.columns]
st.success(f"Loaded **{uploaded.name}**" + (f" \u00B7 sheet **{sheet_name}**" if sheet_name else "")
           + f" \u2014 {len(df)} rows, {len(df.columns)} columns")

# ---- 2. variable types ----------------------------------------------------
st.markdown("### 2. Variable types")
st.caption("The app guesses numerical vs. categorical. Change **Type** where the guess is wrong "
           "(e.g. a group coded 1/2/3 should be *categorical*).")

dataset_key = f"{uploaded.name}::{sheet_name}"
if st.session_state.get("_anova_file") != dataset_key:
    meta = {}
    for col in df.columns:
        vtype, n, uniq = detect_type(df[col])
        meta[col] = {"type": "auto", "detected": vtype, "n": n, "missing": len(df) - n, "unique": uniq}
    st.session_state["anova_meta"] = meta
    st.session_state["_anova_file"] = dataset_key
    st.session_state.pop("anova_result", None)
meta = st.session_state["anova_meta"]

editor_df = pd.DataFrame([{
    "Variable": c, "Type": meta[c]["type"], "Auto-detected": meta[c]["detected"].capitalize(),
    "n (non-missing)": meta[c]["n"], "n (missing)": meta[c]["missing"], "Unique values": meta[c]["unique"],
} for c in df.columns])

edited = st.data_editor(
    editor_df,
    column_config={
        "Type": st.column_config.SelectboxColumn(options=["auto", "numerical", "categorical"], required=True),
        "Variable": st.column_config.TextColumn(disabled=True),
        "Auto-detected": st.column_config.TextColumn(disabled=True),
        "n (non-missing)": st.column_config.NumberColumn(disabled=True),
        "n (missing)": st.column_config.NumberColumn(disabled=True),
        "Unique values": st.column_config.NumberColumn(disabled=True),
    },
    hide_index=True, use_container_width=True, key="anova_var_editor",
)
for _, row in edited.iterrows():
    meta[row["Variable"]]["type"] = row["Type"]

numeric_cols = [c for c in df.columns if effective_type(meta[c]) == "numerical"]
categorical_cols = [c for c in df.columns if effective_type(meta[c]) == "categorical"]

# ---- 3. variables & options -------------------------------------------------
st.markdown("### 3. Outcome(s), grouping variable(s) and tests")
c1, c2 = st.columns(2)
with c1:
    outcomes = st.multiselect("Outcome variable(s) (numerical)", options=numeric_cols)
with c2:
    group_cols = st.multiselect("Grouping variable(s) (categorical)", options=categorical_cols,
                                help="Each grouping variable gets its own one-way analysis.")

o1, o2, o3, o4 = st.columns([1.6, 1.6, 1.6, 1])
with o1:
    test_mode = st.radio("Test selection", ["auto", "parametric", "nonparametric"],
                         format_func=lambda x: {"auto": "Auto (normality-based)",
                                                "parametric": "Parametric (ANOVA)",
                                                "nonparametric": "Non-parametric (Kruskal\u2013Wallis)"}[x])
    display = st.radio("Descriptive statistics", ["auto", "mean_sd", "median_iqr", "both"],
                       format_func=lambda x: {"auto": "Auto (normality-based)", "mean_sd": "Mean \u00B1 SD",
                                              "median_iqr": "Median (IQR)", "both": "Both"}[x])
with o2:
    variance = st.selectbox("ANOVA variant (parametric)", ["auto", "classic", "welch"],
                            format_func=lambda x: {"auto": "Auto (Levene decides)", "classic": "Classic (equal variances)",
                                                   "welch": "Welch (unequal variances)"}[x])
    posthoc_param = st.selectbox("Post hoc \u2014 parametric", ["auto", "tukey", "gameshowell", "bonferroni"],
                                 format_func=lambda x: {"auto": "Auto (Tukey, or Games-Howell if Welch)",
                                                        "tukey": "Tukey HSD", "gameshowell": "Games-Howell",
                                                        "bonferroni": "Bonferroni t-tests"}[x])
with o3:
    posthoc_np = st.selectbox("Post hoc \u2014 non-parametric (Mann-Whitney p adjustment)",
                              ["holm", "bonferroni", "bh", "none"],
                              format_func=lambda x: {"holm": "Holm", "bonferroni": "Bonferroni",
                                                     "bh": "Benjamini-Hochberg (FDR)", "none": "None"}[x])
    posthoc_only_sig = st.checkbox("Post hoc only when overall test is significant", value=False)
with o4:
    alpha = st.number_input("Significance level (\u03B1)", 0.001, 0.5, 0.05, 0.01)
    decimals = st.number_input("Decimal places", 0, 6, 2, 1)

ready = True
if not outcomes:
    st.warning("Select at least one outcome variable.")
    ready = False
if not group_cols:
    st.warning("Select at least one grouping variable.")
    ready = False
overlap = set(outcomes) & set(group_cols)
if overlap:
    st.error(f"A variable cannot be both outcome and grouping: {', '.join(overlap)}")
    ready = False

if ready:
    for g in group_cols:
        levels = df[g].dropna().astype(str).str.strip()
        counts = levels.value_counts().sort_index()
        st.info(f"Groups in **{g}**: " + ", ".join(f"{k} (n={v})" for k, v in counts.items()))

if st.button("Run analysis", type="primary", disabled=not ready):
    cfg = {"alpha": alpha, "decimals": int(decimals), "test_mode": test_mode, "display": display,
           "variance": variance, "posthoc_param": posthoc_param, "posthoc_np": posthoc_np,
           "posthoc_only_sig": posthoc_only_sig}
    results, messages, flags = {}, [], set()
    for g in group_cols:
        all_levels = sorted(df[g].dropna().astype(str).str.strip().unique())
        all_levels = [x for x in all_levels if x not in ("", "nan")]
        s_rows, a_rows, p_rows = [], [], []
        for o in outcomes:
            res, note = analyze(df, o, g, cfg)
            if res is None:
                messages.append(note)
                continue
            messages.extend(note)
            row = {"Outcome": res["label"]}
            for lv in all_levels:
                row[lv] = res["cells"].get(lv, "—")
            row.update({"Test": res["test"]["name"], "Statistic": res["test"]["stat"],
                        "p-value": fmt_p(res["test"]["p"]), "Effect size": res["test"]["effect"]})
            s_rows.append(row)
            a_rows.extend(res["assumptions"])
            for r in res["posthoc"]:
                p_rows.append({k: r[k] for k in ["Outcome", "Comparison", "Method", "Difference", "CI",
                                                  "Statistic", "p (raw)", "p (adj.)"]})
            flags.add("param" if res["use_param"] else "nonparam")
        if s_rows:
            ci = f"{100 * (1 - alpha):g}% CI"
            post_df = pd.DataFrame(p_rows).rename(columns={"CI": ci}) if p_rows else pd.DataFrame()
            results[g] = {"summary": pd.DataFrame(s_rows), "assumptions": pd.DataFrame(a_rows), "posthoc": post_df}
    st.session_state["anova_result"] = (results, messages, flags, cfg)

if "anova_result" in st.session_state:
    results, messages, flags, cfg = st.session_state["anova_result"]
    alpha_r = cfg["alpha"]

    for m in messages:
        st.warning(m)
    if not results:
        st.stop()

    for g, r in results.items():
        st.markdown(f"## Comparison by **{g}**")
        st.markdown("#### Summary")
        st.markdown(render_html(r["summary"], alpha_r), unsafe_allow_html=True)
        copy_button(r["summary"], "Copy summary table")

        st.markdown("#### Post-hoc pairwise comparisons")
        if r["posthoc"].empty:
            st.caption("No post-hoc table: needs \u22653 groups (with 2 groups the omnibus test is the "
                       "pairwise test) or the overall test was not significant while that filter is on.")
        else:
            st.markdown(render_html(r["posthoc"], alpha_r), unsafe_allow_html=True)
            copy_button(r["posthoc"], "Copy post-hoc table")

        with st.expander("Assumption checks (normality & equal variances)"):
            st.markdown(render_html(r["assumptions"], alpha_r), unsafe_allow_html=True)
            copy_button(r["assumptions"], "Copy assumption checks")

    footnotes = [
        "Values are mean \u00B1 SD or median (IQR) as indicated in the Outcome column.",
        f"Bold p-values indicate statistical significance at \u03B1={alpha_r}.",
        "Normality assessed per group with Shapiro-Wilk; equal variances with Levene's test (median-centred).",
    ]
    if "param" in flags:
        footnotes.append("Parametric: one-way ANOVA (Welch's ANOVA if variances unequal). Post-hoc differences "
                         "are mean differences (first group \u2212 second group); Tukey/Games-Howell p-values and "
                         "CIs are familywise-adjusted; Bonferroni CIs use \u03B1/m.")
    if "nonparam" in flags:
        footnotes.append("Non-parametric: Kruskal\u2013Wallis H with effect size \u03B5\u00B2=H/(N\u22121). Post hoc: pairwise "
                         "Mann-Whitney U; difference is the Hodges-Lehmann median difference (first \u2212 second) with a "
                         "distribution-free CI (widened to \u03B1/m under Holm/Bonferroni).")
    st.caption("  \n".join(footnotes))

    d1, d2 = st.columns(2)
    with d1:
        st.download_button("Download Excel", build_excel(results), file_name="anova_results.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    with d2:
        st.download_button("Download Word", build_docx(results, alpha_r, footnotes), file_name="anova_results.docx",
                           mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
